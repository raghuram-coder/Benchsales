"""BenchPilot v1 API server. FastAPI + vanilla JS frontend."""
import base64
import io
import json
import os
import secrets
import threading
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, UploadFile, Request
from fastapi.responses import JSONResponse
from fastapi.staticfiles import StaticFiles

from app import db, emptype, portals, usa
from app import matcher as matcher_mod
from app import tailor as tailor_mod
from app.skills import extract_skills
from app.sources import adzuna, jsearch, remoteok, remotive, arbeitnow, dice, urlimport

BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = BASE_DIR / "frontend"

SOURCES = [adzuna, jsearch, remoteok, remotive, arbeitnow, dice]

app = FastAPI(title="BenchPilot", version="1.0")


@app.middleware("http")
async def auth_gate(request: Request, call_next):
    pwd = os.environ.get("BENCHPILOT_PASSWORD")
    if pwd:
        ok = False
        auth = request.headers.get("authorization", "")
        if auth.startswith("Basic "):
            try:
                decoded = base64.b64decode(auth[6:]).decode("utf-8", "ignore")
                _, _, given = decoded.partition(":")
                ok = secrets.compare_digest(given, pwd)
            except Exception:
                ok = False
        if not ok:
            return JSONResponse({"detail": "authentication required"},
                                status_code=401,
                                headers={"WWW-Authenticate": 'Basic realm="BenchPilot"'})
    return await call_next(request)


@app.on_event("startup")
def _startup():
    db.init_db()
    t = threading.Thread(target=_auto_collect_loop, daemon=True,
                         name="benchpilot-auto-collect")
    t.start()


# ---------- auto-refresh scheduler ----------
_collect_lock = threading.Lock()
_collect_status = {
    "last_run_at": None,      # ISO timestamp of last finished run
    "last_result": None,      # {"jobs_new": n, "matches_new": n}
    "next_run_at": None,      # ISO timestamp of next scheduled run
    "running": False,
}


def _do_collect() -> dict:
    """Run one collection cycle (manual or scheduled). Thread-safe."""
    from collector import run
    with _collect_lock:
        _collect_status["running"] = True
        try:
            summary = run()
        finally:
            _collect_status["running"] = False
    _collect_status["last_run_at"] = datetime.now(timezone.utc).isoformat()
    _collect_status["last_result"] = {
        "jobs_new": summary.get("jobs_new", 0),
        "matches_new": summary.get("matches_new", 0),
        "errors": summary.get("errors", []),
    }
    return summary


def _auto_collect_loop():
    """Background loop: re-pull enabled sources on the configured schedule."""
    while True:
        try:
            interval = int(db.get_setting("collect_interval_minutes", "60") or 0)
        except (ValueError, TypeError):
            interval = 0
        if interval <= 0:
            _collect_status["next_run_at"] = None
            time.sleep(60)
            continue
        _collect_status["next_run_at"] = (
            datetime.now(timezone.utc) + timedelta(minutes=interval)).isoformat()
        time.sleep(interval * 60)
        try:
            _do_collect()
        except Exception:
            pass  # collector records per-source errors; keep the loop alive


@app.get("/api/collect/status")
def api_collect_status():
    try:
        interval = int(db.get_setting("collect_interval_minutes", "60") or 0)
    except (ValueError, TypeError):
        interval = 0
    return {
        "interval_minutes": interval,
        "running": _collect_status["running"],
        "last_run_at": _collect_status["last_run_at"],
        "next_run_at": _collect_status["next_run_at"] if interval > 0 else None,
        "last_result": _collect_status["last_result"],
    }


# ---------- consultants ----------
@app.get("/api/consultants")
def api_list_consultants():
    out = []
    for c in db.list_consultants():
        r = db.get_resume(c["id"])
        c["has_resume"] = r is not None
        c["skills"] = json.loads(r["skills_json"]) if r else []
        c["resume_filename"] = r["filename"] if r else ""
        out.append(c)
    return out


@app.post("/api/consultants")
def api_create_consultant(data: dict):
    if not (data.get("name") or "").strip():
        raise HTTPException(400, "name is required")
    return db.create_consultant(data)


@app.get("/api/consultants/{cid}")
def api_get_consultant(cid: int):
    c = db.get_consultant(cid)
    if not c:
        raise HTTPException(404, "consultant not found")
    r = db.get_resume(cid)
    c["skills"] = json.loads(r["skills_json"]) if r else []
    c["resume_filename"] = r["filename"] if r else ""
    c["resume_text"] = r["raw_text"] if r else ""
    return c


@app.put("/api/consultants/{cid}")
def api_update_consultant(cid: int, data: dict):
    c = db.update_consultant(cid, data)
    if not c:
        raise HTTPException(404, "consultant not found")
    return c


@app.delete("/api/consultants/{cid}")
def api_delete_consultant(cid: int):
    if not db.delete_consultant(cid):
        raise HTTPException(404, "consultant not found")
    return {"ok": True}


def _parse_resume(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if ext == "txt":
        return data.decode("utf-8", "ignore")
    raise HTTPException(400, "unsupported file type (use .pdf, .docx, or .txt)")


@app.post("/api/consultants/{cid}/resume")
async def api_upload_resume(cid: int, file: UploadFile = File(...)):
    if not db.get_consultant(cid):
        raise HTTPException(404, "consultant not found")
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(400, "file too large (10 MB max)")
    text = _parse_resume(file.filename or "resume", data).strip()
    if not text:
        raise HTTPException(400, "could not extract text from file")
    skills = extract_skills(text)
    row = db.upsert_resume(cid, file.filename or "resume", text, skills)
    return {"filename": row["filename"], "skills": skills,
            "chars": len(text)}


# ---------- jobs ----------
@app.get("/api/jobs")
def api_list_jobs(source: str = "", q: str = "", limit: int = 50,
                  emp: str = "", include_unspecified: bool = False):
    """emp = comma list of full-time/c2c/w2/... tags (OR). Jobs whose type
    could not be detected are hidden when emp is set, unless
    include_unspecified=true."""
    return db.list_jobs(source=source, q=q, limit=limit,
                        emp=emptype.parse_wanted(emp),
                        include_unspecified=include_unspecified)


@app.get("/api/emp-types")
def api_emp_types():
    return [{"tag": t, "label": emptype.TAG_LABELS[t]} for t in emptype.TAGS]


@app.get("/api/portals")
def api_portals(title: str = "", location: str = "", emp: str = "",
                days: int = 7):
    """Pre-filled search links for Dice, Indeed, LinkedIn, ZipRecruiter..."""
    if not title.strip():
        raise HTTPException(400, "title is required")
    return portals.build_links(title, location, emptype.parse_wanted(emp),
                               max(1, min(days, 30)))


@app.get("/api/portals/check")
def api_portals_check():
    """Can the server reach each portal? (Runs ~8 quick HTTP requests.)"""
    return portals.check_reachability()


@app.post("/api/jobs/import-url")
def api_import_url(data: dict):
    url = (data.get("url") or "").strip()
    if not url:
        raise HTTPException(400, "url is required")
    try:
        job = urlimport.import_url(url)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"could not import URL: {e}")
    jid = db.insert_job(job)
    job["id"] = jid
    job["duplicate"] = jid is None
    if jid:
        matcher_mod.run_all()  # score consultants against the new posting
    return job


@app.post("/api/collect")
def api_collect():
    return _do_collect()


LIVE_SOURCES = [adzuna, jsearch, remoteok, remotive, arbeitnow]  # dice: no public API


@app.post("/api/jobs/live-search")
def api_live_search(data: dict):
    """Search the boards right now for a job title (+ optional location).

    Returns listings without saving them; the frontend saves explicitly
    via /api/jobs/save-live.
    """
    title = (data.get("title") or "").strip()
    if not title:
        raise HTTPException(400, "title is required")
    location = (data.get("location") or "").strip()
    settings = db.get_settings()
    try:
        enabled_sources = set(json.loads(settings.get("enabled_sources") or "[]"))
    except json.JSONDecodeError:
        enabled_sources = set(db.DEFAULT_ENABLED_SOURCES)
    query = {"title": title, "location": location}
    emp_wanted = emptype.parse_wanted(",".join(data.get("emp") or [])
                                      if isinstance(data.get("emp"), list)
                                      else (data.get("emp") or ""))
    include_unspecified = bool(data.get("include_unspecified"))
    usa_only = (settings.get("usa_only", "1") == "1") and data.get("usa_only", True)
    # let the live search steer the per-source filters the same way collection does
    if emp_wanted:
        settings = {**settings, "collect_emp_types": ",".join(emp_wanted)}
    jobs: list[dict] = []
    sources: dict = {}
    for mod in LIVE_SOURCES:
        name = mod.NAME
        label = getattr(mod, "LABEL", name)
        if name not in enabled_sources:
            sources[name] = {"label": label, "status": "disabled in Settings", "count": 0}
            continue
        if not mod.enabled(settings):
            sources[name] = {"label": label, "status": "needs API key — add it in Settings", "count": 0}
            continue
        try:
            fetched = mod.fetch(query, settings)
            if usa_only:
                fetched = usa.filter_us(fetched)
            for j in fetched:
                j["temp_id"] = f"{name}:{j.get('source_id')}"
                j["emp_tags"] = emptype.classify(
                    j.get("employment_type", ""), j.get("title", ""),
                    j.get("description", ""))
            if emp_wanted:
                fetched = [j for j in fetched if emptype.matches_filter(
                    j["emp_tags"], emp_wanted, include_unspecified)]
            fetched = fetched[:30]
            jobs.extend(fetched)
            sources[name] = {"label": label, "status": "ok", "count": len(fetched)}
        except Exception as e:  # noqa: BLE001 - one bad source shouldn't kill the search
            sources[name] = {"label": label, "status": f"error: {e}", "count": 0}
    # rank: jobs whose title covers more of the query words come first
    words = title.lower().split()

    def _relevance(job: dict):
        t = (job.get("title") or "").lower()
        c = (job.get("company") or "").lower()
        hit_title = sum(1 for w in words if w in t)
        hit_all = sum(1 for w in words if w in t or w in c)
        return (hit_title / len(words), hit_all / len(words))

    jobs.sort(key=_relevance, reverse=True)
    return {"jobs": jobs, "sources": sources}


@app.post("/api/jobs/save-live")
def api_save_live(data: dict):
    """Save live-search listings into the job board (dedupes) and re-run matching."""
    jobs = data.get("jobs") or []
    if not isinstance(jobs, list) or not jobs:
        raise HTTPException(400, "jobs list is required")
    saved, dupes = 0, 0
    for j in jobs:
        j.pop("temp_id", None)
        j.pop("emp_tags", None)  # recomputed on insert
        if db.insert_job(j):
            saved += 1
        else:
            dupes += 1
    matches_new = matcher_mod.run_all()
    return {"saved": saved, "duplicates": dupes, "matches_new": matches_new}


# ---------- matches ----------
@app.get("/api/matches")
def api_list_matches(consultant_id: int = 0, min_score: float = 0):
    return db.list_matches(consultant_id or None, min_score)


@app.get("/api/matches/{mid}")
def api_get_match(mid: int):
    m = db.get_match(mid)
    if not m:
        raise HTTPException(404, "match not found")
    return m


@app.post("/api/matches/{mid}/tailor")
def api_tailor(mid: int):
    m = db.get_match(mid)
    if not m:
        raise HTTPException(404, "match not found")
    settings = db.get_settings()
    result = tailor_mod.tailor_match(m, settings)
    tid = db.insert_tailored(mid, result["text"], result["tailored_by"],
                             result["added_skills_flagged"])
    return {"id": tid, "tailored_by": result["tailored_by"],
            "added_skills_flagged": result["added_skills_flagged"]}


@app.get("/api/tailored/{tid}")
def api_get_tailored(tid: int):
    t = db.get_tailored(tid)
    if not t:
        raise HTTPException(404, "not found")
    m = db.get_match(t["match_id"])
    if m:
        t["original_text"] = m.get("resume_text", "")
    return t


@app.post("/api/matches/{mid}/queue")
def api_queue(mid: int, data: dict | None = None):
    if not db.get_match(mid):
        raise HTTPException(404, "match not found")
    tid = (data or {}).get("tailored_resume_id")
    return db.queue_application(mid, tid)


# ---------- applications ----------
@app.get("/api/applications")
def api_list_applications(status: str = ""):
    return db.list_applications(status)


@app.patch("/api/applications/{aid}")
def api_update_application(aid: int, data: dict):
    try:
        a = db.update_application(aid, data.get("status"), data.get("notes"))
    except ValueError as e:
        raise HTTPException(400, str(e))
    if not a:
        raise HTTPException(404, "application not found")
    return a


# ---------- settings / sources ----------
MASKED_KEYS = {"adzuna_app_key", "rapidapi_key", "llm_api_key"}


def _mask(key: str, value: str) -> str:
    if key in MASKED_KEYS and value:
        return (value[:3] + "***" + value[-2:]) if len(value) > 6 else "***"
    return value


@app.get("/api/settings")
def api_get_settings():
    return {k: _mask(k, v) for k, v in db.get_settings().items()}


@app.put("/api/settings")
def api_put_settings(data: dict):
    allowed = {"match_threshold", "search_queries", "enabled_sources",
               "adzuna_app_id", "adzuna_app_key", "rapidapi_key",
               "llm_base_url", "llm_api_key", "llm_model",
               "collect_interval_minutes", "usa_only", "collect_emp_types"}
    for k, v in data.items():
        if k not in allowed:
            raise HTTPException(400, f"unknown setting: {k}")
        if k in MASKED_KEYS and isinstance(v, str) and "***" in v:
            continue  # masked value echoed back -> keep existing
        db.set_setting(k, v if isinstance(v, str) else json.dumps(v))
    return api_get_settings()


@app.get("/api/sources/status")
def api_sources_status():
    settings = db.get_settings()
    try:
        enabled_sources = set(json.loads(settings.get("enabled_sources") or "[]"))
    except json.JSONDecodeError:
        enabled_sources = set(db.DEFAULT_ENABLED_SOURCES)
    out = []
    for mod in SOURCES:
        name = mod.NAME
        last = {}
        try:
            last = json.loads(settings.get(f"last_run_{name}") or "{}")
        except json.JSONDecodeError:
            pass
        configured = mod.enabled(settings)
        state = "disabled"
        if name == "dice":
            state = "link-out only (no public API) — use Portal search links, Adzuna or URL import"
        elif name in enabled_sources:
            state = "ready" if configured else "needs credentials"
        out.append({
            "name": name,
            "label": getattr(mod, "LABEL", name),
            "enabled": name in enabled_sources and configured,
            "configured": configured,
            "state": state,
            "last_run": last,
        })
    out.append({"name": "urlimport", "label": urlimport.LABEL,
                "enabled": True, "configured": True,
                "state": "manual (Import URL box)", "last_run": {}})
    return out


if FRONTEND_DIR.exists():
    app.mount("/", StaticFiles(directory=str(FRONTEND_DIR), html=True),
              name="frontend")
