"""Core logic tests (stdlib unittest; no network, no web framework needed).

Run from the project root:   python -m unittest discover -s tests -v
"""
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app import autolearn, db, emptype, matcher, portals, skills, usa  # noqa: E402
from app.sources import urlimport  # noqa: E402


class EmpTypeTests(unittest.TestCase):
    def check(self, et, title, desc, expected):
        self.assertEqual(set(emptype.classify(et, title, desc)), set(expected),
                         f"{et!r} {title!r} {desc!r}")

    def test_aggregator_type_words(self):
        self.check("full_time", "Java Developer", "", ["fulltime"])
        self.check("CONTRACTOR", "QA Engineer", "", ["contract"])
        self.check("part_time", "Dev", "", ["parttime"])

    def test_c2c_w2_in_description(self):
        self.check("", "Java Dev", "Open to C2C or W2. Hybrid in Dallas, TX",
                   ["c2c", "w2", "contract"])
        self.check("", "Dev", "Corp-to-Corp, 1099 and W2 fine, 12 month contract",
                   ["c2c", "w2", "1099", "contract"])
        self.check("", "Sr Python Dev (C2C)", "", ["c2c", "contract"])

    def test_negations(self):
        self.check("", "Dev", "W2 only. No C2C. No third party.", ["w2"])
        self.check("", "Dev", "C2C not accepted. W2 candidates only", ["w2"])
        self.check("", "Dev", "C2C is not allowed", [])
        self.check("", "Dev", "No corp to corp please; W2 contract role",
                   ["w2", "contract"])

    def test_real_posting_no_third_party_list(self):
        # real Adzuna/Columbus posting that was wrongly tagged C2C + 1099
        desc = ("$55/hr W2 Contract. USC/GC/GC EAD Only- W2 Requirement\n\n"
                "Note: No third party/C2C or 1099\n\nJob Overview")
        self.check("contract", "Java Developer", desc, ["w2", "contract"])
        self.check("", "Dev", "Open to C2C or W2, no 1099", ["c2c", "w2", "contract"])
        self.check("", "Dev", "No sponsorship, C2C ok", ["c2c", "contract"])
        self.check("", "Dev", "W2 required", ["w2"])

    def test_contract_to_hire(self):
        self.check("", "Data Eng", "Contract to hire, W2", ["w2", "contract", "c2h"])

    def test_boilerplate_is_not_a_type(self):
        self.check("", "Data Eng", "Great benefits for full-time employees", [])
        self.check("", "Contracts Manager", "manage vendor agreements", [])

    def test_filter_semantics(self):
        self.assertTrue(emptype.matches_filter(["c2c"], ["c2c", "w2"]))
        self.assertFalse(emptype.matches_filter(["fulltime"], ["c2c", "w2"]))
        self.assertFalse(emptype.matches_filter([], ["c2c"]))
        self.assertTrue(emptype.matches_filter([], ["c2c"], include_unspecified=True))
        self.assertTrue(emptype.matches_filter(["parttime"], []))


class UsaTests(unittest.TestCase):
    def test_locations(self):
        yes = ["Houston, TX", "Remote", "Worldwide", "", "Austin, Texas",
               "Remote - USA", "Americas", "Dallas-Fort Worth, TX",
               "Remote (Canada or US)", "Vienna, VA", "Melbourne, FL",
               "Boston (remote)", "Europe, USA, Canada, APAC (remote)"]
        no = ["Berlin, Germany", "Europe", "London, UK", "Toronto, Ontario, Canada",
              "Remote, India", "Bengaluru, Karnataka, India",
              "Bishkek, Bishkek City, Kyrgyzstan (remote)",
              "Kuala Lumpur, Malaysia (remote)", "Seoul (remote)",
              "Budapest, (remote)", "Melbourne (remote)",
              "Uluberia-II, (remote)"]
        for loc in yes:
            self.assertTrue(usa.is_us_job({"location": loc}), loc)
        for loc in no:
            self.assertFalse(usa.is_us_job({"location": loc}), loc)


class PortalTests(unittest.TestCase):
    def test_links(self):
        links = {l["id"]: l["url"] for l in
                 portals.build_links("Java Developer", "Houston, TX",
                                     ["fulltime", "c2c", "w2"], 7)}
        for pid in ("dice", "indeed", "linkedin", "ziprecruiter", "glassdoor"):
            self.assertIn(pid, links)
            self.assertTrue(links[pid].startswith("https://"))
        self.assertIn("Java+Developer", links["dice"])
        self.assertIn("filters.employmentType=THIRD_PARTY", links["dice"])
        self.assertIn("f_JT=F%2CC", links["linkedin"])
        self.assertIn("C2C", links["indeed"])


class DbTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._old = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "t.db"
        db.init_db()

    def tearDown(self):
        db.DB_PATH = self._old
        self.tmp.cleanup()

    def job(self, n, **kw):
        j = {"source": "t", "source_id": f"id{n}", "title": f"Java Dev {n}",
             "company": f"Co{n}", "location": "Houston, TX", "remote_flag": 0,
             "url": "", "description": "Java Spring Boot", "posted_at": "",
             "salary": "", "employment_type": ""}
        j.update(kw)
        return j

    def test_tags_stored_and_filtered(self):
        db.insert_job(self.job(1, employment_type="full_time"))
        db.insert_job(self.job(2, description="C2C or W2 contract"))
        db.insert_job(self.job(3, description="W2 only. No C2C."))
        db.insert_job(self.job(4))  # nothing detectable
        titles = lambda **kw: sorted(j["title"] for j in db.list_jobs(**kw))
        self.assertEqual(titles(emp=["c2c"]), ["Java Dev 2"])
        self.assertEqual(titles(emp=["w2"]), ["Java Dev 2", "Java Dev 3"])
        self.assertEqual(titles(emp=["fulltime", "c2c"]), ["Java Dev 1", "Java Dev 2"])
        self.assertEqual(titles(emp=["c2c"], include_unspecified=True),
                         ["Java Dev 2", "Java Dev 4"])
        self.assertEqual(len(db.list_jobs()), 4)
        self.assertIn("c2c", db.list_jobs(emp=["c2c"])[0]["emp_tags"])

    def test_stale_tags_are_retagged_on_start(self):
        db.insert_job(self.job(1, description="W2 Requirement. Note: No third party/C2C or 1099"))
        with db.get_conn() as c:   # simulate tags saved by the old buggy rules
            c.execute("UPDATE jobs SET emp_tags=',c2c,w2,1099,'")
            c.execute("DELETE FROM settings WHERE key='emp_tags_version'")
            c.commit()
        db.init_db()
        self.assertEqual(db.list_jobs()[0]["emp_tags"], ["w2"])

    def test_dedupe_uses_location(self):
        self.assertTrue(db.insert_job(self.job(1, title="Java Dev", company="Acme")))
        # same title+company+location from another source -> duplicate
        self.assertIsNone(db.insert_job(self.job(2, title="java dev", company="ACME")))
        # same title+company, different city -> a separate opening
        self.assertTrue(db.insert_job(self.job(3, title="Java Dev", company="Acme",
                                               location="Dallas, TX")))

    def test_migration_from_old_schema(self):
        p = Path(self.tmp.name) / "old.db"
        con = sqlite3.connect(p)
        con.executescript("""
        CREATE TABLE consultants (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
          email TEXT DEFAULT '', phone TEXT DEFAULT '', location TEXT DEFAULT '',
          visa_status TEXT DEFAULT '', linkedin_url TEXT DEFAULT '', notes TEXT DEFAULT '',
          created_at TEXT NOT NULL);
        CREATE TABLE jobs (id INTEGER PRIMARY KEY AUTOINCREMENT, source TEXT NOT NULL,
          source_id TEXT NOT NULL, title TEXT DEFAULT '', company TEXT DEFAULT '',
          location TEXT DEFAULT '', remote_flag INTEGER DEFAULT 0, url TEXT DEFAULT '',
          description TEXT DEFAULT '', posted_at TEXT DEFAULT '', salary TEXT DEFAULT '',
          employment_type TEXT DEFAULT '', fetched_at TEXT NOT NULL, UNIQUE(source, source_id));
        INSERT INTO jobs(source, source_id, title, description, fetched_at)
          VALUES ('x','1','Old job','Corp to corp welcome','2026-01-01');
        """)
        con.commit()
        con.close()
        db.DB_PATH = p
        db.init_db()
        db.init_db()  # idempotent
        j = db.list_jobs(emp=["c2c"])
        self.assertEqual([x["title"] for x in j], ["Old job"])
        c = db.create_consultant({"name": "A", "emp_pref": "w2, c2c ,bogus"})
        self.assertEqual(c["emp_pref"], "w2,c2c")

    def test_matcher_respects_consultant_preference(self):
        w2_only = {"emp_pref": "w2"}
        self.assertTrue(matcher.emp_compatible(w2_only, ["w2", "c2c"]))
        self.assertFalse(matcher.emp_compatible(w2_only, ["c2c"]))
        self.assertTrue(matcher.emp_compatible(w2_only, []))      # unknown passes
        self.assertTrue(matcher.emp_compatible({"emp_pref": ""}, ["c2c"]))  # no pref

    def test_run_all_end_to_end(self):
        c1 = db.create_consultant({"name": "W2 Person", "emp_pref": "w2"})
        c2 = db.create_consultant({"name": "Any Person"})
        txt = "Senior Java Developer\nJava Spring Boot microservices REST API AWS Docker PostgreSQL"
        for c in (c1, c2):
            db.upsert_resume(c["id"], "r.txt", txt, matcher.extract_skills(txt))
        desc = ("Requirements:\n- Java, Spring Boot, microservices, REST API\n"
                "- AWS, Docker, PostgreSQL\n")
        db.insert_job(self.job(1, title="Senior Java Developer",
                               description=desc + "C2C only"))
        db.insert_job(self.job(2, title="Senior Java Developer", company="Other",
                               description=desc + "W2 only"))
        n = matcher.run_all(threshold=10)
        by_c = {}
        for m in db.list_matches():
            by_c.setdefault(m["consultant_name"], []).append(m["emp_tags"])
        self.assertEqual(len(by_c["Any Person"]), 2)
        self.assertEqual(by_c["W2 Person"], [["w2"]])  # C2C-only job skipped
        self.assertEqual(n, 3)
        self.assertEqual(matcher.run_all(threshold=10), 0)  # idempotent


    def test_run_all_for_one_consultant(self):
        c1 = db.create_consultant({"name": "A"})
        c2 = db.create_consultant({"name": "B"})
        txt = "Senior Java Developer\nJava Spring Boot microservices REST API AWS Docker"
        for c in (c1, c2):
            db.upsert_resume(c["id"], "r.txt", txt, matcher.extract_skills(txt))
        db.insert_job(self.job(1, description="Requirements:\n- Java, Spring Boot, "
                               "microservices, REST API, AWS, Docker\n"))
        self.assertEqual(matcher.run_all(threshold=10, consultant_id=c1["id"]), 1)
        self.assertEqual([m["consultant_name"] for m in db.list_matches()], ["A"])
        self.assertEqual(matcher.run_all(threshold=10, consultant_id=c1["id"]), 0)


class LocationScoreTests(unittest.TestCase):
    def sim(self, a, b):
        return matcher._jaccard(matcher._loc_tokens(a), matcher._loc_tokens(b))

    def test_state_abbreviation_matches_state_name(self):
        self.assertGreater(self.sim("California, USA", "Irvine, CA"), 0)
        self.assertEqual(self.sim("Irvine, CA", "Irvine, CA"), 1.0)
        self.assertEqual(self.sim("Houston, TX", "Irvine, CA"), 0.0)
        self.assertGreater(self.sim("New York, NY", "New York City"), 0.5)


class ResumeFileTests(unittest.TestCase):
    def make_docx(self):
        import io
        from docx import Document
        d = Document()
        d.add_paragraph("JANE DOE")
        d.add_paragraph("TECHNICAL SKILLS")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Languages:"
        t.cell(0, 1).text = "Java, Python"
        t.cell(1, 0).text = "Tools:"
        t.cell(1, 1).text = "Jenkins, Docker"
        d.add_paragraph("EXPERIENCE")
        d.add_paragraph("Built things at Acme")
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    def test_table_stays_where_it_is(self):
        from app import resume
        text = resume.parse("jane.docx", self.make_docx())
        lines = text.splitlines()
        self.assertEqual(lines[:2], ["JANE DOE", "TECHNICAL SKILLS"])
        self.assertEqual(lines[2], "Languages: Java, Python")
        self.assertEqual(lines[3], "Tools: Jenkins, Docker")
        self.assertEqual(lines[4], "EXPERIENCE")

    def test_txt_and_unsupported(self):
        from app import resume
        self.assertEqual(resume.parse("a.txt", b"hello"), "hello")
        with self.assertRaises(ValueError):
            resume.parse("a.exe", b"MZ")
        with self.assertRaises(Exception):
            resume.parse("bad.docx", b"not really a docx")


class TailorTests(unittest.TestCase):
    RESUME = ("JANE DOE\nQA Engineer\nSUMMARY\nQA person.\n\nTECHNICAL SKILLS\n"
              "Languages: Java, Python, SQL\n"
              "Tools: Jenkins, Docker, RestAssured, Selenium\n\n"
              "EXPERIENCE\nBuilt things.")

    def test_categorised_skills_keep_their_lines(self):
        from app import tailor
        skills = matcher.extract_skills(self.RESUME)
        out = tailor.keyword_tailor(self.RESUME, skills,
                                    "Requirements: RestAssured, Selenium, Docker")
        lines = out.splitlines()
        self.assertIn("Languages: Java, Python, SQL", lines)
        # JD skills move to the front of their own line, nothing is dropped
        self.assertIn("Tools: Docker, RestAssured, Selenium, Jenkins", lines)
        self.assertFalse(any(ln.startswith("Skills:") for ln in lines))

    def test_plain_skills_list_is_still_reordered(self):
        from app import tailor
        txt = "JANE DOE\nSKILLS\nJava, Python, Docker\n\nEXPERIENCE\nx"
        out = tailor.keyword_tailor(txt, ["java", "python", "docker"],
                                    "Requirements: Docker")
        self.assertIn("Skills: docker, java, python", out)


class SkillsTests(unittest.TestCase):
    def test_qa_and_ai_skills_found(self):
        from app import skills
        got = set(skills.extract_skills(
            "Frameworks: Rest Assured, Selenium WebDriver, SoapUI, Allure "
            "Reporting. AI: Claude API integration, Generative AI QA, "
            "ISTQB CT-AI."))
        for s in ("restassured", "selenium webdriver", "soapui", "allure",
                  "claude api", "generative ai", "istqb"):
            self.assertIn(s, got)

    def test_teams_word_is_not_microsoft_teams(self):
        from app import skills
        self.assertNotIn("microsoft teams", skills.extract_skills(
            "Led cross-functional teams of 14 engineers"))
        self.assertIn("microsoft teams", skills.extract_skills(
            "Daily standups on MS Teams"))


class AutoTests(unittest.TestCase):
    """The 'works without humans' parts: learning, queries, quota, rescoring."""
    RESUME = ("Sharan Murali\nQA Automation Engineer | Dallas, TX\n\n"
              "SUMMARY\nTester.\n\nTECHNICAL SKILLS\n"
              "Languages: Java, Python, SQL\n"
              "Test Mgmt: Zephyr Scale, Xray (Jira), TestRail\n"
              "Tools: Selenium, Postman, Jenkins, and, etc\n\n"
              "EXPERIENCE\nWorked with Cobol Mainframe daily.\n")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._old = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "t.db"
        db.init_db()
        skills.set_learned([])

    def tearDown(self):
        skills.set_learned([])
        matcher._jd_skills.cache_clear()
        db.DB_PATH = self._old
        self.tmp.cleanup()

    def test_candidates_come_only_from_skills_section(self):
        got = autolearn.candidate_skills(self.RESUME)
        self.assertIn("zephyr scale", got)
        self.assertIn("xray", got)
        self.assertIn("testrail", got)
        for bad in ("and", "etc", "cobol mainframe", "java", "selenium"):
            self.assertNotIn(bad, got)

    def test_learning_changes_extraction_and_is_removable(self):
        c = db.create_consultant({"name": "S", "location": "Dallas, TX"})
        db.upsert_resume(c["id"], "r.txt", self.RESUME, skills.extract_skills(self.RESUME))
        self.assertNotIn("zephyr scale", db.get_resume(c["id"])["skills_json"])
        out = autolearn.maintain()
        self.assertIn("zephyr scale", out["learned"])
        self.assertEqual(out["resumes_refreshed"], 1)
        self.assertIn("zephyr scale", db.get_resume(c["id"])["skills_json"])
        self.assertIn("zephyr scale", skills.extract_skills("need Zephyr Scale"))
        self.assertEqual(autolearn.maintain()["learned"], [])  # nothing new 2nd time
        self.assertTrue(db.delete_learned_skill("zephyr scale"))
        autolearn.block_skill("zephyr scale")
        autolearn.maintain()   # must not learn it again from the same resume
        self.assertNotIn("zephyr scale", skills.extract_skills("need Zephyr Scale"))

    def test_rescoring_improves_scores_and_never_deletes(self):
        c = db.create_consultant({"name": "S"})
        db.upsert_resume(c["id"], "r.txt", self.RESUME, skills.extract_skills(self.RESUME))
        desc = ("Requirements:\n- Java, Selenium, Zephyr Scale, TestRail, Xray\n")
        jid = db.insert_job({"source": "t", "source_id": "1", "title": "QA Automation Engineer",
                             "company": "Co", "location": "Dallas, TX", "remote_flag": 0,
                             "url": "", "description": desc, "posted_at": "",
                             "salary": "", "employment_type": ""})
        self.assertEqual(matcher.run_all(threshold=0), 1)
        before = db.list_matches()[0]["score"]
        autolearn.maintain()  # learns the three tools -> job needs them, resume has them
        after = db.list_matches()
        self.assertEqual(len(after), 1)
        self.assertGreater(after[0]["score"], before)

    def test_queries_built_from_resumes(self):
        c1 = {"raw_text": self.RESUME, "skills": ["java"], "location": "Dallas, TX"}
        c2 = {"raw_text": "Ravi K\nJava Full Stack Developer\nJava Spring", "skills": ["java"],
              "location": ""}
        qs = autolearn.build_queries([c1, c2], [{"title": "Manual Role", "location": "Austin"}], 8)
        titles = [(q["title"], q["location"]) for q in qs]
        self.assertEqual(titles[0], ("QA Automation Engineer", "Dallas, TX"))
        self.assertIn(("Java Full Stack Developer", "Remote"), titles)
        self.assertIn(("Manual Role", "Austin"), titles)
        self.assertEqual(len(titles), len(set(titles)))
        self.assertEqual(len(autolearn.build_queries([c1, c2], [], 2)), 2)
        # nobody uploaded anything yet -> fall back to the typed queries
        self.assertEqual(autolearn.queries_for_run({"auto_queries": "1"}),
                         db.DEFAULT_SEARCH_QUERIES)

    def test_interval_respects_adzuna_daily_budget(self):
        base = {"collect_interval_minutes": "15", "enabled_sources": '["adzuna"]',
                "auto_queries": "0", "adzuna_daily_budget": "80"}
        # no keys -> no slowdown
        self.assertEqual(autolearn.effective_interval_minutes(base), 15)
        keyed = {**base, "adzuna_app_id": "a", "adzuna_app_key": "b"}
        n = len(autolearn.queries_for_run(keyed))        # 4 default queries -> 8 calls/run
        runs = 80 // (n * 2)
        self.assertEqual(autolearn.effective_interval_minutes(keyed), -(-1440 // runs))
        self.assertLessEqual(runs * n * 2, 80)
        self.assertEqual(autolearn.effective_interval_minutes({**keyed, "collect_interval_minutes": "0"}), 0)

    def test_collector_stops_at_daily_budget(self):
        import collector
        db.set_setting("adzuna_daily_budget", "6")
        qs = [{"title": f"T{i}", "location": ""} for i in range(5)]
        s = {"budget_skipped": {}}
        self.assertEqual(len(collector._within_budget(qs, s)), 3)   # 3 x 2 calls = 6
        self.assertEqual(s["budget_skipped"]["adzuna"], 2)
        db.usage_add("adzuna", 6)
        self.assertEqual(collector._within_budget(qs, {"budget_skipped": {}}), [])


class SsrfTests(unittest.TestCase):
    def test_private_hosts_blocked(self):
        for u in ["http://127.0.0.1:8741/api/settings",
                  "http://169.254.169.254/latest/meta-data/",
                  "http://localhost/", "file:///etc/passwd",
                  "http://10.0.0.5/", "http://[::1]/"]:
            with self.assertRaises(ValueError, msg=u):
                urlimport._assert_public(u)


if __name__ == "__main__":
    unittest.main()
